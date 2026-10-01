from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from datetime import date as D
H=["ID","Nombre / dirección","Comuna","Rol SII","Tipo","m² útiles","Fecha de compra","Tipo de financiamiento","Costo total (UF)","Capital propio aportado (UF)","Banco","Monto del crédito (UF)","Tasa anual (%)","Plazo (años)","Fecha primer dividendo","Dividendo mensual (UF)","Saldo insoluto hoy (UF)","Estado","Arriendo mensual (UF)","Inicio del contrato","Fin del contrato","Contribuciones (UF/año)","Gastos comunes no recuperables (UF/año)","Valor comercial (UF)","Valor banco (UF)","Fecha de tasación","Tasador","Archivo PDF tasación","Plusvalía sobre costo (%)","Notas"]
R=[
["P11","Locales comerciales La Portada, Antonia López de Bello 934 al 968","Independencia","288-11 (planilla: 248-11)","Local comercial",518,None,None,None,None,None,None,None,None,None,None,None,"Arrendada",118.6,None,None,None,None,18050,18050,D(2024,11,7),"Perito judicial V. Correa (causa Santander C-15903-2023)","Tasacion Centro Comercial La Portada Independencia (1).pdf",None,"Arriendo = renta neta capitalizada en la tasación, no contrato. Informe judicial $685,5 MM a UF nov 2024. Valor libro $365.000.000."],
["P05","Terrazas 401, Av. del Parque 5275 Of. 401 + 7 estac.","Huechuraba","6047-20, 6047-234, 235, 236, 237, 245, 256, 257","Oficina",283.02,None,None,None,None,"Banco de Chile",12296.05,4.94,15,D(2025,1,1),99.91,None,"Arrendada",77.5,None,None,None,None,16635,16635,D(2024,5,28),"CGDV para Banco de Chile","Tasacion Oficina 401 Terrazas.pdf",None,"Arrendatario Importadora Hevia SpA. Arriendo según tasación. Crédito op. 899.601 de 12.100 UF; el monto incluye intereses capitalizados de 4 meses de gracia. Valor libro $630.702.717."],
["P03","Terrazas 202, Av. del Parque 5275 Of. 202 + 5 estac. + bodega 46","Huechuraba","6047-11, 6047-136 a 140, 6047-298","Oficina",275,None,None,None,None,"Itaú",9761.98,4.48,12,D(2026,1,20),87.75,None,"Arrendada",85,None,None,None,None,14100,14100,D(2025,6,12),"N. Latorre para Itaú","Tasacion Itau Terrazas 202 (1).pdf",None,"Arrendatario Corporación Vial S.A. Arriendo según tasación. Tasa implícita desde la tabla de desarrollo. Valor libro $552.179.500."],
["P09","Oficina Nueva York 25, piso 3","Santiago","00029-00003","Oficina",311,None,None,None,None,"Santander",3990,4.94,12,D(2026,7,8),36.79,None,"En reparación",None,None,None,None,None,8225,8225,D(2026,1,20),"Propiteq (estimación online)","Tasacion Oficina Nueva York 25.pdf",None,"Recibida 15/07, remodelando. Valor es estimación automática. Tasa implícita."],
["P07","Ñuñoa, Los Talaveras 120 depto 405 + estac. 79 + bodega 44","Ñuñoa","3950-678, 3950-768, 3950-834","Departamento",99,None,None,None,None,None,None,None,None,None,None,None,None,None,None,None,None,None,8030,8030,D(2026,8,5),"BancoEstado","Tasacion Ñuñoa.pdf",None,"PDF: $328.000.000. Arrendatario sin informar. Valor libro $220.500.000."],
["P01","Moneda 812, Of. 705","Santiago","00025-00291","Oficina",297,None,None,None,None,None,None,None,None,None,None,None,"Arrendada",54,None,None,None,None,9800,7840,D(2025,7,25),"N. Latorre para Itaú (TMA251378)","TMA251378 (3).pdf",None,"Arrendatario Vision Optics Chile SpA. Arriendo según tasación. Otra tasación: BancoEstado sep 2025, 6.509 UF (TASACION Moneda.pdf). Valor libro $384.252.660."],
["P10","Huérfanos 1294, Of. 51","Santiago","00107-00223","Oficina",203,None,None,None,None,None,None,None,None,None,None,None,"En reparación",46.7,None,None,None,None,7657,7657,D(2026,7,17),"Propiteq (estimación online)","Tasacion Oficina Huerfanos.pdf",None,"Recibida junio, remodelando. Arriendo = estimación Propiteq. Valor libro $124.585.059."],
["P04","Work Center Miraflores, Av. Ventisquero 2225 Módulo 42 Nave D","Renca","02899-00296","Bodega",None,None,None,None,None,"Banco de Chile",5000,4.5,12,D(2026,2,1),46.72,None,"Arrendada",None,None,None,None,None,7545,7545,D(2025,11,3),"Banco de Chile (correo)","Tasacion Work Center Miraflores.pdf",None,"Arrendatario Medio Mundo SpA. Crédito op. 166.411. Valor libro $299.747.458."],
["P06","Europa 401, Av. del Parque 4680 + estac. 113, 114 + bodega 68","Huechuraba","6061-38, 6061-109, 6061-110, 6061-73","Oficina",126,None,None,None,None,"BancoEstado",3971.25,3.97,12,D(2023,9,30),34.58,None,"Arrendada",None,None,None,None,None,5367,5367,D(2023,5,10),"BancoEstado","Tasacion Oficial Edificio Europa 401.pdf",None,"Arrendatario Mendes Holler Ingeniería SpA. Crédito op. 116424450. Valor libro $164.000.000."],
["P02","Patio Mayor, Av. del Parque 4980 Of. 331 + estac. 28, 29","Huechuraba","6055-76, 6055-187, 6055-188, 6055-136","Oficina",111,None,None,None,None,"Itaú",3552,4.58,12,D(2025,5,5),32.1,None,"Arrendada",32,None,None,None,None,4440,4440,D(2024,12,26),"N. Latorre para Itaú","TMA242294 - TASACION OFICINA Patio Mayor (1).pdf",None,"Arrendatario Más a Más SpA. Arriendo según tasación. Tasa implícita. Valor libro $170.384.300."],
["P08","Recoleta, Edificio Ilumina, Av. Valdivieso 411, deptos 25, 27, 85 y 105","Recoleta","2371-26, 2371-82, 2371-104, 2371-106, 2371-28, 2371-68, 2371-108","Departamento",255.74,None,None,None,None,None,None,None,None,None,None,None,"En compra",57.85,None,None,None,None,18200,15268,D(2026,9,4),"Roberto Nieto para Banco de Chile","TASACION DEPTO 25, 27, 85, 105 (4 archivos)",None,"10% pagado ($40.129.495), deuda banco aprox. $360 MM por confirmar. Arriendo estimado en tasaciones. Nuevos sin uso, recepción final pendiente."],
]
col={h:i+1 for i,h in enumerate(H)}
from openpyxl.utils import get_column_letter as L
c=lambda h:L(col[h])
wb=Workbook();ws=wb.active;ws.title="Propiedades"
ws.append(H)
hf=Font(name="Arial",bold=True,color="FFFFFF");fill=PatternFill("solid",fgColor="1F3A5F")
for i in range(1,len(H)+1):
    x=ws.cell(1,i);x.font=hf;x.fill=fill;x.alignment=Alignment(wrap_text=True,vertical="center")
N=20
for r in range(2,N+1):
    row=R[r-2] if r-2<len(R) else [None]*len(H)
    for i,v in enumerate(row,1):
        if v is not None: ws.cell(r,i,v)
    Lc,Mc,Nc,Oc=c("Monto del crédito (UF)"),c("Tasa anual (%)"),c("Plazo (años)"),c("Fecha primer dividendo")
    i_=f"(1+{Mc}{r}/1200)";n_=f"({Nc}{r}*12)"
    ws.cell(r,col["Saldo insoluto hoy (UF)"],f'=IF(OR({Lc}{r}="",{Mc}{r}="",{Nc}{r}="",{Oc}{r}=""),"",IF(TODAY()<{Oc}{r},{Lc}{r},ROUND({Lc}{r}*({i_}^{n_}-{i_}^MIN({n_},DATEDIF({Oc}{r},TODAY(),"M")+1))/({i_}^{n_}-1),1)))')
    X,I=c("Valor comercial (UF)"),c("Costo total (UF)")
    ws.cell(r,col["Plusvalía sobre costo (%)"],f'=IF(OR({X}{r}="",N({I}{r})=0),"",ROUND(({X}{r}/{I}{r}-1)*100,1))')
    for h in ("Fecha de compra","Fecha primer dividendo","Inicio del contrato","Fin del contrato","Fecha de tasación"):
        ws.cell(r,col[h]).number_format="dd-mm-yyyy"
Hc,Jc=c("Tipo de financiamiento"),c("Capital propio aportado (UF)")
dv=DataValidation(type="list",formula1='"Leaseback,Compra con banco"',allow_blank=True,showErrorMessage=True,errorTitle="Tipo de financiamiento",error="Elige Leaseback o Compra con banco.")
ws.add_data_validation(dv);dv.add(f"{Hc}2:{Hc}{N}")
dv2=DataValidation(type="custom",formula1=f'${Hc}2="Compra con banco"',allow_blank=True,showErrorMessage=True,errorTitle="Capital propio",error="Solo se ingresa capital propio si el financiamiento es Compra con banco.")
ws.add_data_validation(dv2);dv2.add(f"{Jc}2:{Jc}{N}")
dv3=DataValidation(type="list",formula1='"Arrendada,Vacante,En reparación,En venta,Vendida,Uso propio,En compra"',allow_blank=True)
ws.add_data_validation(dv3);c3=c("Estado");dv3.add(f"{c3}2:{c3}{N}")
gray=PatternFill("solid",fgColor="E7E7E7");yel=PatternFill("solid",fgColor="FFF2B3")
ws.conditional_formatting.add(f"{Jc}2:{Jc}{N}",FormulaRule(formula=[f'AND(${Hc}2="Compra con banco",${Jc}2="")'],fill=yel))
ws.conditional_formatting.add(f"{Jc}2:{Jc}{N}",FormulaRule(formula=[f'AND($A2<>"",${Hc}2<>"Compra con banco")'],fill=gray))
w={"ID":7,"Nombre / dirección":40,"Rol SII":24,"Tipo de financiamiento":18,"Tasador":28,"Archivo PDF tasación":30,"Notas":60}
for h,i in col.items(): ws.column_dimensions[L(i)].width=w.get(h,13)
ws.row_dimensions[1].height=45;ws.freeze_panes="C2"
wb.save("registro_v2.xlsx")
