from openpyxl import load_workbook
import datetime as dt, json
wb=load_workbook('detalle.xlsx',data_only=True)
def num(v):
    if v in (None,'','-'): return 0.0
    if isinstance(v,(int,float)): return float(v)
    return float(str(v).replace('.','').replace(',','.'))
def rows(name):
    for r in wb[name].iter_rows(values_only=True):
        yield list(r)
out={}
def add(key,items,first):
    # items: list of (n, pago_uf, capital_uf, interes_uf)
    items.sort()
    d=[]
    for n,pago,cap,intr in items:
        m=first.month-1+int(n)-1
        d.append((dt.date(first.year+m//12,m%12+1,1),pago,cap,intr))
    out[key]=d
# Generic: locate header row then read
def sched(name,ncol,cols,div=1,first=None,skipcheck=None):
    it=[]; hdr=False
    for r in rows(name):
        vals=[v for v in r]
        if not hdr:
            if any(isinstance(v,str) and v.strip() in ('FCH.VCT','F.Vcto','N° dividendos') for v in vals): hdr=True
            continue
        n=vals[ncol]
        if not isinstance(n,(int,float)) or n<1: continue
        g=lambda k: num(vals[k])/div if k is not None else 0
        cap,intr,pago=g(cols[0]),g(cols[1]),g(cols[2]) if cols[2] is not None else 0
        if not pago: pago=cap+intr
        it.append((int(n),pago,cap,intr))
    return it
# Column positions (0-based in the raw row)
def show(name):
    for i,r in enumerate(rows(name)):
        if i in (9,10,11,12): print(name,list(enumerate(r))[:12])
for n in wb.sheetnames[4:]: show(n)

def rd(name,ncol,cap,intr,pago,div=1):
    it=[]
    for r in rows(name):
        n=r[ncol] if len(r)>ncol else None
        if not isinstance(n,(int,float)) or isinstance(n,bool) or n<1 or n>400 or n!=int(n): continue
        try:
            c=num(r[cap])/div; i=num(r[intr])/div; p=num(r[pago])/div if pago is not None else c+i
        except Exception: continue
        if c==0 and i==0: continue
        it.append((int(n),p,c,i))
    # dedupe by n
    seen={}; [seen.setdefault(x[0],x) for x in it]
    return sorted(seen.values())
S=wb.sheetnames
cfg={
 2:(S[3],3,5,6,4,1,dt.date(2025,5,1),'Itaú',3552.0),
 3:(S[4],3,5,6,4,1,dt.date(2026,1,1),'Itaú',9761.98),
 4:(S[5],1,4,5,8,1,dt.date(2026,2,1),'Banco de Chile',5000.0),
 5:(S[6],1,4,5,8,1,dt.date(2025,1,1),'Banco de Chile',12100.0),
 6:(S[7],2,3,4,6,1,dt.date(2023,9,1),'BancoEstado',3971.25),
 9:(S[8],2,3,4,None,10000,dt.date(2026,7,1),'Santander',3990.0),
}
res={}
for k,(name,nc,c,i,p,dv,first,banco,monto) in cfg.items():
    it=rd(name,nc,c,i,p,dv)
    caps=sum(x[2] for x in it)
    sched=[]
    for n,pago,cap,intr in it:
        m=first.month-1+n-1; sched.append((first.year+m//12,first.month and m%12+1,pago,cap,intr))
    # saldo at 2026-10-01 : monto - capital paid before Oct 2026
    paid=sum(s[3] for s in sched if (s[0],s[1])<(2026,10))
    saldo=monto-paid
    div=[s[2] for s in sched if (s[0],s[1])==(2026,10)]
    fin=sched[-1][:2]
    # yearly from Oct 2026: year t=1 is Oct26-Sep27
    yr={}
    for y,mo,pago,cap,intr in sched:
        if (y,mo)<(2026,10): continue
        t=(y-2026)*12+mo-10; t=t//12+1
        a=yr.setdefault(t,[0,0,0]); a[0]+=pago;a[1]+=cap;a[2]+=intr
    res[k]=dict(banco=banco,monto=monto,cuotas=len(it),suma_cap=round(caps,1),saldo=round(saldo,1),div=round(div[0],2) if div else None,fin=f"{fin[1]:02d}/{fin[0]}",
        anual=[[t,round(v[0],1),round(v[1],1),round(v[2],1)] for t,v in sorted(yr.items())])
    print(k,banco,monto,len(it),round(caps,1),'saldo',round(saldo,1),'div',res[k]['div'],'fin',res[k]['fin'],'yr1',res[k]['anual'][0])
json.dump(res,open('creditos.json','w'))
